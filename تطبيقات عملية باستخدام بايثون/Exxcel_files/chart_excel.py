import openpyxl
from pathlib import Path









excel_sheet=openpyxl.load_workbook(Path.home()/Path('Desktop','employees.xlsx'))
sheet_one=excel_sheet['one_sheet']


'''#chart

title=openpyxl.chart.Reference(sheet_one,min_col=1,max_col=1,min_row=1,max_row=6)
data=openpyxl.chart.Reference(sheet_one,min_col=2,max_col=2,min_row=1,max_row=6)
chart=openpyxl.chart.BarChart()


chart.title='my chart'
chart.add_data(data=data)
chart.set_categories(title)

sheet_one.add_chart(chart,'E8')
excel_sheet.save(filename=Path.home()/Path('Desktop','emplyees.xlsx'))'''


'''

#charts
right=openpyxl.chart.Reference(sheet_one,min_row=2,max_row=7,min_col=1,max_col=1)
bottom=openpyxl.chart.Reference(sheet_one,min_row=2,max_row=7,min_col=3,max_col=3)
chart=openpyxl.chart.BarChart()
chart.title='employees salaries'
chart.add_data(data=bottom)
chart.set_categories(right)
sheet_one.add_chart(chart,'E9')
excel_sheet.save(filename=Path.home()/Path('Desktop','employees.xlsx'))

'''



